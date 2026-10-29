#!/usr/bin/env python3
"""Fill result tables and performance plots in the included DOCX report from benchmark.csv."""
import csv
import sys
from pathlib import Path

from docx import Document
from docx.shared import Inches


def set_cell(cell, value):
    p = cell.paragraphs[0]
    if p.runs:
        p.runs[0].text = str(value)
        for run in p.runs[1:]:
            run.text = ""
    else:
        p.add_run(str(value))


def insert_picture_at_marker(paragraph, marker, path):
    if marker not in paragraph.text:
        return False
    for run in list(paragraph.runs):
        run._element.getparent().remove(run._element)
    path = Path(path)
    if path.exists():
        paragraph.alignment = 1
        paragraph.add_run().add_picture(str(path), width=Inches(6.4))
        return True
    paragraph.add_run(f"График не найден: {path.name}. Сначала выполните scripts/plot_benchmarks.py.")
    return False


def main():
    if len(sys.argv) < 4:
        print("Использование: python scripts/fill_report.py benchmark.csv Отчет_ЛР3.docx Отчет_ЛР3_заполненный.docx")
        return 2
    csv_path, input_docx, output_docx = map(Path, sys.argv[1:4])
    if not csv_path.exists() or not input_docx.exists():
        print("Не найден CSV или DOCX. Проверьте пути.")
        return 1
    with csv_path.open("r", encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
    records = {}
    for row in rows:
        try:
            records[(row["task"], row["backend"], int(row["input_size"]))] = row
        except (KeyError, ValueError):
            pass

    doc = Document(input_docx)
    matrix_sizes = [32, 64, 128, 256]
    for table in doc.tables:
        if not table.rows:
            continue
        first = table.rows[0].cells[0].text.strip()
        if first == "Размер матрицы":
            for row_index, size in enumerate(matrix_sizes, start=1):
                if row_index >= len(table.rows):
                    break
                cpu = records.get(("determinant", "CPU", size))
                gpu = records.get(("determinant", "GPU", size))
                set_cell(table.rows[row_index].cells[0], f"{size} × {size}")
                set_cell(table.rows[row_index].cells[1], f"{float(cpu['time_ms']):.3f}" if cpu else "нет данных")
                set_cell(table.rows[row_index].cells[2], f"{float(gpu['time_ms']):.3f}" if gpu else "GPU не измерен")
                set_cell(table.rows[row_index].cells[3], "да" if cpu and gpu and cpu.get('found') == gpu.get('found') == '1' else ("проверить" if gpu else "—"))
        elif first == "Алгоритм":
            cpu = records.get(("md5", "CPU", 1_000_000)) or next((v for (task, backend, _), v in records.items() if task == "md5" and backend == "CPU"), None)
            gpu = records.get(("md5", "GPU", 1_000_000)) or next((v for (task, backend, _), v in records.items() if task == "md5" and backend == "GPU"), None)
            set_cell(table.rows[1].cells[1], cpu.get("input_size", "—") if cpu else "нет данных")
            set_cell(table.rows[1].cells[2], f"{float(cpu['time_ms']):.3f}" if cpu else "нет данных")
            set_cell(table.rows[1].cells[3], f"{float(gpu['time_ms']):.3f}" if gpu else "GPU не измерен")
            success = bool(cpu and gpu and cpu.get("found") == gpu.get("found") == "1")
            set_cell(table.rows[1].cells[4], "найдено на CPU/GPU" if success else ("только CPU / сверить" if cpu else "нет данных"))

    filled = 0
    for paragraph in doc.paragraphs:
        if "[[MATRIX_CHART]]" in paragraph.text:
            filled += int(insert_picture_at_marker(paragraph, "[[MATRIX_CHART]]", "matrix_benchmark.png"))
        elif "[[MD5_CHART]]" in paragraph.text:
            filled += int(insert_picture_at_marker(paragraph, "[[MD5_CHART]]", "md5_benchmark.png"))
        elif "[[TEST_OUTPUT]]" in paragraph.text:
            for run in list(paragraph.runs):
                run._element.getparent().remove(run._element)
            paragraph.add_run("Вставьте сюда фактический вывод `lab3 test`.")

    doc.save(output_docx)
    print(f"Отчёт сохранён: {output_docx}")
    if filled < 2:
        print("Предупреждение: не все графики были вставлены. Запустите scripts/plot_benchmarks.py из корня проекта.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
